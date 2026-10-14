import datetime as dt
import io
from pathlib import Path

import openpyxl
import pytest

from app.importers.excel_grid import GridFormatError, parse_grid
from app.importers.grid_export import export_grid
from tests.conftest import FIXTURES, MEGA

SAMPLE = FIXTURES / "grid_sample.xlsx"
REAL = Path(__file__).resolve().parents[2] / "Mega 6-45 Vietnam.xlsx"


def test_fixture_parses_like_the_users_sheet():
    r = parse_grid(SAMPLE, MEGA)
    assert [d.draw_no for d in r.draws] == [998, 1000, 1003, 1005, 1006]  # oldest -> newest
    assert r.draws[-1].numbers == [4, 8, 15, 16, 23, 42]
    assert r.draws[-1].draw_date == dt.date(2026, 3, 13)

    # draw number (+ maybe date) but no numbers => missing, never filled in
    assert [(m.draw_no, m.draw_date) for m in r.missing] == [
        (1004, dt.date(2026, 3, 8)),
        (1002, None),
        (1001, None),
    ]
    # 5 numbers, and a cell whose value does not match its column
    assert len(r.rejected) == 2
    reasons = " | ".join(x.reason for x in r.rejected)
    assert "found 5 numbers, expected exactly 6" in reasons
    assert "holds 11 but its column is 10" in reasons


def test_export_round_trips_through_parse():
    parsed = parse_grid(SAMPLE, MEGA)
    blob = export_grid(parsed.draws, MEGA, missing_draw_nos=[m.draw_no for m in parsed.missing])
    again = parse_grid(blob, MEGA)
    assert [(d.draw_no, d.draw_date, d.numbers) for d in again.draws] == [
        (d.draw_no, d.draw_date, d.numbers) for d in parsed.draws
    ]
    assert sorted(m.draw_no for m in again.missing) == [1001, 1002, 1004]
    assert again.rejected == []


def test_export_layout_matches_user_sheet():
    parsed = parse_grid(SAMPLE, MEGA)
    ws = openpyxl.load_workbook(io.BytesIO(export_grid(parsed.draws, MEGA))).active
    assert ws.title == "45"
    assert ws["A3"].value == "Date" and ws["C3"].value == "Draw"
    assert ws["E3"].value == 1 and ws["AW3"].value == 45
    # newest on top, number in its own column (4 -> column H = E+3)
    assert ws["C5"].value == 1006
    assert ws["H5"].value == 4
    assert ws["B5"].value == "Fri"


def test_export_carries_codencode_wordmark_in_two_colours():
    ws = openpyxl.load_workbook(io.BytesIO(export_grid([], MEGA)), rich_text=True).active
    rt = ws["A2"].value
    parts = [(b.text, b.font.color.rgb) for b in rt]
    assert parts == [("codencode", "FF000000"), (".my", "FF00DCB4")]


def test_wrong_layout_is_reported_not_guessed(tmp_path):
    wb = openpyxl.Workbook()
    wb.active["A1"] = "nothing here"
    p = tmp_path / "bad.xlsx"
    wb.save(p)
    with pytest.raises(GridFormatError, match="header row not found"):
        parse_grid(p, MEGA)


def test_wrong_game_size_is_reported(tmp_path):
    # A 45-column grid read as a 6/55 game must fail loudly (header lacks 46..55).
    from app.validation import GameSpec

    with pytest.raises(GridFormatError, match="every number 1..55"):
        parse_grid(SAMPLE, GameSpec("x", 55, 6), sheet="45")


@pytest.mark.skipif(not REAL.exists(), reason="user's workbook not present")
def test_real_user_workbook_facts():
    """Documents what the user's actual file contains (observed 2026-10-13)."""
    r = parse_grid(REAL, MEGA)
    assert r.rejected == []
    assert len(r.draws) == 57
    assert len(r.missing) == 65  # rows with a draw number but no numbers (1571 and 1504-1542 ...)
    nos = [d.draw_no for d in r.draws]
    assert min(nos) == 1453 and max(nos) == 1574
    assert {d.draw_date.weekday() for d in r.draws} <= {2, 4, 6}  # Wed, Fri, Sun
