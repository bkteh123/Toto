import csv
import io

import openpyxl
import pytest
from sqlalchemy import select

from app.importers.excel_grid import parse_grid
from app.importers.text_csv import parse_text
from app.models import Draw, Game
from tests.conftest import MEGA
from tests.fakesite import load_seed


@pytest.fixture
def loaded(session):
    g = session.scalar(select(Game).where(Game.code == "vn_mega_645"))
    for d in load_seed():
        session.add(Draw(game_id=g.id, draw_no=d.draw_no, draw_date=d.date, numbers=d.numbers, source="t"))
    session.commit()


def _rows(resp):
    assert resp.content.startswith(b"\xef\xbb\xbf"), "CSV downloads carry a BOM so Excel reads UTF-8"
    return list(csv.reader(io.StringIO(resp.content.decode("utf-8-sig"))))


def _sheet(resp):
    assert resp.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    return openpyxl.load_workbook(io.BytesIO(resp.content), data_only=True).active


def test_health_checks_the_database(client):
    r = client.get("/health")
    assert r.status_code == 200 and r.json() == {"status": "ok"}


# ---------- stats ----------


def test_stats_csv_has_one_row_per_number(client, loaded):
    r = client.get("/games/vn_mega_645/export/stats.csv?period=20")
    assert r.status_code == 200
    assert 'filename="vn_mega_645-scores-20.csv"' in r.headers["content-disposition"]
    rows = _rows(r)
    assert rows[0][:3] == ["number", "appearances", "frequency"]
    assert [int(x[0]) for x in rows[1:]] == list(range(1, 46))
    # Appearances in the last 20 draws sum to 20 draws x 6 numbers.
    assert sum(int(x[1]) for x in rows[1:]) == 120
    assert {x[8] for x in rows[1:]} <= {"Hot", "Cold", "Neutral"}


def test_stats_csv_matches_the_stats_endpoint(client, loaded):
    api = client.get("/games/vn_mega_645/stats?period=50&mix=70").json()["rows"]
    rows = _rows(client.get("/games/vn_mega_645/export/stats.csv?period=50&mix=70"))[1:]
    for a, row in zip(api, rows):
        assert int(row[0]) == a["number"] and int(row[1]) == a["appearances"]
        assert float(row[7]) == pytest.approx(a["combined"], abs=1e-4)


def test_stats_xlsx_carries_wordmark_and_disclaimer(client, loaded):
    ws = _sheet(client.get("/games/vn_mega_645/export/stats.xlsx"))
    assert "not probabilities of winning" in ws["A3"].value
    assert "codencode" in str(ws["A2"].value)
    assert ws.cell(5, 1).value == "number" and ws.cell(6, 1).value == 1 and ws.max_row == 5 + 45


def test_stats_export_without_draws_is_a_409_not_an_empty_file(client):
    assert client.get("/games/vn_mega_645/export/stats.csv").status_code == 409


@pytest.mark.parametrize("path", ["stats.pdf", "draws.json", "nope.csv"])
def test_unknown_format_or_kind_is_rejected(client, loaded, path):
    assert client.get(f"/games/vn_mega_645/export/{path}").status_code in (404, 422)


def test_unknown_game_is_404(client):
    assert client.get("/games/nope/export/stats.csv").status_code == 404


# ---------- draws ----------


def test_draws_csv_round_trips_through_the_text_importer(client, loaded):
    r = client.get("/games/vn_mega_645/export/draws.csv")
    assert r.status_code == 200
    text = r.content.decode("utf-8-sig")
    parsed = parse_text(text, MEGA, skip_leading=1)
    seed = load_seed()
    assert not parsed.rejected and len(parsed.draws) == len(seed) == 122
    assert {(d.draw_no, d.draw_date, tuple(d.numbers)) for d in parsed.draws} == {
        (s.draw_no, s.date, tuple(s.numbers)) for s in seed
    }


def test_draws_xlsx_round_trips_through_the_grid_importer(client, loaded):
    r = client.get("/games/vn_mega_645/export/draws.xlsx")
    assert r.status_code == 200
    parsed = parse_grid(r.content, MEGA)
    seed = load_seed()
    assert not parsed.rejected and len(parsed.draws) == len(seed)
    assert {(d.draw_no, tuple(d.numbers)) for d in parsed.draws} == {(s.draw_no, tuple(s.numbers)) for s in seed}


# ---------- combos ----------


def test_combos_csv_echoes_validated_picks(client, loaded):
    body = {"method": "hot", "combinations": [[3, 1, 20, 33, 40, 45], [2, 4, 6, 8, 10, 12]]}
    r = client.post("/games/vn_mega_645/export/combos.csv", json=body)
    assert r.status_code == 200
    assert _rows(r) == [
        ["#", "method", "n1", "n2", "n3", "n4", "n5", "n6"],
        ["1", "hot", "1", "3", "20", "33", "40", "45"],  # sorted
        ["2", "hot", "2", "4", "6", "8", "10", "12"],
    ]


def test_combos_xlsx(client, loaded):
    r = client.post(
        "/games/vn_mega_645/export/combos.xlsx",
        json={"method": "random", "combinations": [[1, 2, 3, 4, 5, 6]]},
    )
    ws = _sheet(r)
    assert [c.value for c in ws[6]][:8] == [1, "random", 1, 2, 3, 4, 5, 6]
    assert "baseline" in ws["A3"].value


@pytest.mark.parametrize(
    "combos",
    [[[1, 2, 3, 4, 5]], [[1, 2, 3, 4, 5, 46]], [[1, 1, 2, 3, 4, 5]], [[0, 1, 2, 3, 4, 5]], []],
)
def test_combos_rejects_illegal_picks(client, loaded, combos):
    r = client.post("/games/vn_mega_645/export/combos.csv", json={"method": "hot", "combinations": combos})
    assert r.status_code == 422


# ---------- bao ----------


def test_bao_csv_matches_the_bao_endpoint(client, loaded):
    api = client.post("/games/vn_mega_645/bao", json={"size": 8, "window": 10}).json()
    r = client.get("/games/vn_mega_645/export/bao.csv?size=8&window=10")
    assert r.status_code == 200
    rows = _rows(r)
    assert rows[0] == ["number", "appearances", "status"]
    assert [int(x[0]) for x in rows[1:]] == api["numbers"]
    assert [int(x[1]) for x in rows[1:]] == [api["counts"][str(n)] for n in api["numbers"]]
    sure = {int(x[0]) for x in rows[1:] if x[2] == "sure_in"}
    assert sure == set(api["sure_in"])


def test_bao_xlsx_has_cost_summary(client, loaded):
    ws = _sheet(client.get("/games/vn_mega_645/export/bao.xlsx?size=5&window=10"))
    labels = {ws.cell(r, 1).value: ws.cell(r, 2).value for r in range(1, ws.max_row + 1)}
    assert labels["lines"] == 40  # Bao 5 on a 6/45 game: 45 - 5 lines
    assert labels["total cost (VND)"] == 400000


def test_bao_export_rejects_invalid_size(client, loaded):
    assert client.get("/games/vn_mega_645/export/bao.csv?size=6&window=10").status_code == 422
    assert client.get("/games/vn_mega_645/export/bao.csv?window=10").status_code == 422
