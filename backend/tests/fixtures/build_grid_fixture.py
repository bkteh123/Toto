"""Builds tests/fixtures/grid_sample.xlsx: a tiny workbook shaped like Mega_6-45_Vietnam.xlsx.

The draws are invented for testing; they are NOT real results.
Run:  python tests/fixtures/build_grid_fixture.py
"""

import datetime as dt
from pathlib import Path

import openpyxl

OUT = Path(__file__).with_name("grid_sample.xlsx")

# (draw_no, date|None, numbers|None, note) newest first, as in the real sheet.
ROWS = [
    (1006, dt.date(2026, 3, 13), [4, 8, 15, 16, 23, 42], "ok"),
    (1005, dt.date(2026, 3, 11), [2, 9, 17, 23, 38, 45], "ok"),
    (1004, dt.date(2026, 3, 8), None, "date but no numbers (like draw 1571 in the real file)"),
    (1003, dt.date(2026, 3, 6), [5, 11, 12, 26, 31, 40], "ok"),
    (1002, None, None, "draw number only"),
    (1001, None, None, "draw number only"),
    (1000, dt.date(2026, 2, 27), [1, 7, 14, 22, 33, 41], "ok"),
    (999, dt.date(2026, 2, 25), [3, 6, 19, 28, 44], "only 5 numbers -> rejected"),
    (998, dt.date(2026, 2, 22), [10, 20, 30, 40, 43, 45], "ok"),
]


def build(path: Path = OUT) -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "45"
    ws["A1"] = "6/45   (MEGA) fixture"
    ws.cell(3, 1, "Date")
    ws.cell(3, 3, "Draw ")  # trailing space exactly like the real header
    for n in range(1, 46):
        ws.cell(3, 4 + n, n)  # E3 = 1 ... AW3 = 45
    r = 5
    for draw_no, date, numbers, _ in ROWS:
        if date:
            ws.cell(r, 1, dt.datetime(date.year, date.month, date.day))
            ws.cell(r, 2, date.strftime("%a"))
        ws.cell(r, 3, draw_no)
        for n in numbers or []:
            ws.cell(r, 4 + n, n)
        r += 1
    # A row whose cell value does not match its column (column 10 holds 11): must be rejected.
    ws.cell(r, 1, dt.datetime(2026, 2, 20))
    ws.cell(r, 3, 997)
    for n in (2, 5, 8, 12, 25):
        ws.cell(r, 4 + n, n)
    ws.cell(r, 4 + 10, 11)  # wrong value in column "10"
    wb.save(path)
    return path


if __name__ == "__main__":
    print(build())
