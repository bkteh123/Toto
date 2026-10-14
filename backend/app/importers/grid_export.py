"""Export draws back to the Excel grid layout so the user can keep working in Excel."""

from __future__ import annotations

import datetime as dt
import io
from collections.abc import Iterable, Sequence

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.branding import stamp_wordmark
from app.importers.common import ParsedDraw
from app.validation import GameSpec

HEADER_ROW = 3
FIRST_DATA_ROW = 5
FIRST_NUMBER_COL = 5  # column E

GOLD = PatternFill("solid", fgColor="F2C230")
LIGHT_GOLD = PatternFill("solid", fgColor="FFF2B3")


def export_grid(
    draws: Sequence[ParsedDraw],
    spec: GameSpec,
    missing_draw_nos: Iterable[int] = (),
    sheet_title: str | None = None,
) -> bytes:
    """Return an .xlsx whose layout `parse_grid` reads back identically.

    Newest draws first. `missing_draw_nos` become rows holding only the draw number, as in
    the user's own sheet, so gaps stay visible.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_title or str(spec.n)

    ws["A1"] = f"{spec.k}/{spec.n}   {spec.name}".strip()
    ws["A1"].font = Font(bold=True, size=14)
    stamp_wordmark(ws, "A2")

    ws.cell(HEADER_ROW, 1, "Date")
    ws.cell(HEADER_ROW, 2, "Day")
    ws.cell(HEADER_ROW, 3, "Draw")
    for n in range(1, spec.n + 1):
        ws.cell(HEADER_ROW, FIRST_NUMBER_COL + n - 1, n)
    for c in range(1, FIRST_NUMBER_COL + spec.n):
        cell = ws.cell(HEADER_ROW, c)
        cell.font = Font(bold=True, color="FF140F06")
        cell.fill = GOLD
        cell.alignment = Alignment(horizontal="center")

    # Row model: (draw_no, date, numbers|None)
    rows: list[tuple[int | None, dt.date | None, list[int] | None]] = [
        (d.draw_no, d.draw_date, list(d.numbers)) for d in draws
    ]
    present = {r[0] for r in rows if r[0] is not None}
    rows += [(m, None, None) for m in missing_draw_nos if m not in present]

    if all(r[0] is not None for r in rows):
        rows.sort(key=lambda r: r[0], reverse=True)
    else:
        rows.sort(key=lambda r: (r[1] or dt.date.min, r[0] or 0), reverse=True)

    for i, (draw_no, date, numbers) in enumerate(rows):
        r = FIRST_DATA_ROW + i
        if date is not None:
            ws.cell(r, 1, date).number_format = "dd/mm/yyyy"
            ws.cell(r, 2, date.strftime("%a"))
        if draw_no is not None:
            ws.cell(r, 3, draw_no)
        for n in numbers or ():
            cell = ws.cell(r, FIRST_NUMBER_COL + n - 1, n)
            cell.fill = LIGHT_GOLD
            cell.alignment = Alignment(horizontal="center")

    ws.column_dimensions["A"].width = 12
    ws.column_dimensions["B"].width = 6
    ws.column_dimensions["C"].width = 7
    for n in range(1, spec.n + 1):
        ws.column_dimensions[get_column_letter(FIRST_NUMBER_COL + n - 1)].width = 4
    ws.freeze_panes = ws.cell(FIRST_DATA_ROW, FIRST_NUMBER_COL)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
