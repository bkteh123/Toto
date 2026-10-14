"""Import the user's Excel "grid" format (e.g. Mega_6-45_Vietnam.xlsx).

Layout: header row has `Date` in column A, `Draw` in column C and the numbers 1..N in
consecutive columns from E. Each data row has the date, weekday and draw number, and a
drawn number sits in its own column as its own value. Newest draws are at the top.
"""

from __future__ import annotations

import datetime as dt
import io
import re
from pathlib import Path
from typing import Any, BinaryIO

import openpyxl
from openpyxl.utils import get_column_letter

from app.importers.common import (
    DateParseError,
    MissingRow,
    ParsedDraw,
    ParseResult,
    RejectedRow,
    find_date,
)
from app.validation import GameSpec

HEADER_SEARCH_ROWS = 15


class GridFormatError(ValueError):
    """The workbook does not look like the expected grid layout."""


def _blank(v: Any) -> bool:
    return v is None or (isinstance(v, str) and not v.strip())


def _as_int(v: Any) -> int | None:
    """Explicit, lossless integer reading of a cell: 7, 7.0 and "7" -> 7; anything else None."""
    if isinstance(v, bool):
        return None
    if isinstance(v, int):
        return v
    if isinstance(v, float) and v.is_integer():
        return int(v)
    if isinstance(v, str) and re.fullmatch(r"\s*\d+\s*", v):
        return int(v)
    return None


def _as_date(v: Any) -> dt.date | None:
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    if isinstance(v, str):
        d, _ = find_date(v.strip())  # may raise DateParseError
        return d
    return None


def _pick_sheet(wb: openpyxl.Workbook, spec: GameSpec, sheet: str | None):
    if sheet is not None:
        if sheet not in wb.sheetnames:
            raise GridFormatError(f"sheet {sheet!r} not found; workbook has {wb.sheetnames}")
        return wb[sheet]
    if str(spec.n) in wb.sheetnames:
        return wb[str(spec.n)]
    return wb.worksheets[0]


def parse_grid(
    source: str | Path | bytes | BinaryIO, spec: GameSpec, sheet: str | None = None
) -> ParseResult:
    if isinstance(source, bytes):
        source = io.BytesIO(source)
    wb = openpyxl.load_workbook(source, data_only=True)
    ws = _pick_sheet(wb, spec, sheet)

    # Locate the header row: `Date` in column A within the first rows.
    header_row = None
    for r in range(1, min(ws.max_row, HEADER_SEARCH_ROWS) + 1):
        a = ws.cell(r, 1).value
        if isinstance(a, str) and a.strip().lower() == "date":
            header_row = r
            break
    if header_row is None:
        raise GridFormatError(
            f"header row not found: expected 'Date' in column A within the first "
            f"{HEADER_SEARCH_ROWS} rows of sheet {ws.title!r}"
        )

    draw_col = None
    number_cols: dict[int, int] = {}  # column index -> number
    for c in range(1, ws.max_column + 1):
        v = ws.cell(header_row, c).value
        if isinstance(v, str) and v.strip().lower() == "draw" and draw_col is None:
            draw_col = c
        n = _as_int(v) if not isinstance(v, str) else None
        if n is not None and 1 <= n <= spec.n:
            number_cols[c] = n
    if draw_col is None:
        raise GridFormatError("header row has no 'Draw' column")
    if sorted(number_cols.values()) != list(range(1, spec.n + 1)):
        raise GridFormatError(
            f"header row must contain every number 1..{spec.n} exactly once; "
            f"found {len(number_cols)} numeric header cells"
        )

    result = ParseResult(ordering="date")
    for r in range(header_row + 1, ws.max_row + 1):
        date_raw = ws.cell(r, 1).value
        draw_raw = ws.cell(r, draw_col).value
        cells = {
            c: ws.cell(r, c).value for c in number_cols if not _blank(ws.cell(r, c).value)
        }
        if _blank(date_raw) and _blank(draw_raw) and not cells:
            continue

        row_text = f"row {r}"
        draw_no = None if _blank(draw_raw) else _as_int(draw_raw)
        try:
            draw_date = None if _blank(date_raw) else _as_date(date_raw)
        except DateParseError as exc:
            result.rejected.append(RejectedRow(r, row_text, str(exc)))
            continue
        if not _blank(draw_raw) and draw_no is None:
            result.rejected.append(RejectedRow(r, row_text, f"draw number {draw_raw!r} is not an integer"))
            continue
        if not _blank(date_raw) and draw_date is None:
            result.rejected.append(RejectedRow(r, row_text, f"date {date_raw!r} is not a date"))
            continue

        if not cells:
            # Draw number (and maybe a date) but no numbers yet: report, never fill in.
            result.missing.append(MissingRow(r, draw_no, draw_date))
            continue

        problems: list[str] = []
        picked: list[int] = []
        for c, v in cells.items():
            n = _as_int(v)
            col_number = number_cols[c]
            if n is None or n != col_number:
                problems.append(f"{get_column_letter(c)}{r} holds {v!r} but its column is {col_number}")
            else:
                picked.append(n)
        if problems:
            result.rejected.append(RejectedRow(r, row_text, "; ".join(problems)))
            continue
        if len(picked) != spec.k:
            result.rejected.append(
                RejectedRow(r, row_text, f"found {len(picked)} numbers, expected exactly {spec.k}")
            )
            continue
        if draw_date is None and draw_no is None:
            result.rejected.append(RejectedRow(r, row_text, "row has numbers but no date or draw number"))
            continue
        result.draws.append(
            ParsedDraw(numbers=sorted(picked), draw_date=draw_date, draw_no=draw_no, line_no=r)
        )

    result.draws.sort(key=lambda d: (d.draw_date or dt.date.min, d.draw_no or 0))
    return result
