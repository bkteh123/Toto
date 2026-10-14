"""File downloads: CSV (UTF-8 with BOM, opens cleanly in Excel) and Excel workbooks.

stats, draws and bao are deterministic functions of the stored draws and the query
parameters, so they are GET downloads. Generated combinations are random and cannot be
recomputed, so the combos export is a POST that echoes the picks the user is looking at
(after validating them against the game's rules).
"""

from __future__ import annotations

import csv
import io
from typing import Literal
from urllib.parse import quote

import openpyxl
from fastapi import APIRouter, Depends, HTTPException, Query, Response
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from pydantic import BaseModel, Field
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.analysis import DEFAULT_WEIGHTS
from app.analysis.generator import MAX_COUNT
from app.api.picks import BaoRequest, bao
from app.api.routes import _game
from app.api.stats import NOTE, game_stats
from app.branding import stamp_wordmark
from app.db import get_session
from app.importers.common import ParsedDraw
from app.importers.grid_export import export_grid
from app.importers.service import missing_draw_numbers
from app.models import Draw
from app.validation import DrawValidationError, GameSpec, validate_numbers

router = APIRouter()

Fmt = Literal["csv", "xlsx"]
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
CSV = "text/csv; charset=utf-8"
GOLD = PatternFill("solid", fgColor="F2C230")
HEADER_ROW = 5  # rows 1-3: title, wordmark, note; row 4 is blank


def _attachment(data: bytes, fmt: Fmt, stem: str) -> Response:
    name = f"{stem}.{fmt}"
    return Response(
        content=data,
        media_type=XLSX if fmt == "xlsx" else CSV,
        headers={
            "Content-Disposition": f"attachment; filename=\"{name}\"; filename*=UTF-8''{quote(name)}",
            "Cache-Control": "no-store",
        },
    )


def _csv(header: list[str], rows: list[list]) -> bytes:
    buf = io.StringIO(newline="")
    w = csv.writer(buf, lineterminator="\n")
    w.writerow(header)
    w.writerows(["" if c is None else c for c in r] for r in rows)
    return b"\xef\xbb\xbf" + buf.getvalue().encode("utf-8")


def _xlsx(
    title: str,
    note: str,
    header: list[str],
    rows: list[list],
    sheet: str = "Export",
    footer: list[list] | None = None,
) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet[:31]
    ws["A1"] = title
    ws["A1"].font = Font(bold=True, size=14)
    stamp_wordmark(ws, "A2")
    ws["A3"] = note
    ws["A3"].font = Font(italic=True, color="FF5A4A1A")
    for c, h in enumerate(header, start=1):
        cell = ws.cell(HEADER_ROW, c, h)
        cell.font = Font(bold=True, color="FF140F06")
        cell.fill = GOLD
        cell.alignment = Alignment(horizontal="center")
        ws.column_dimensions[get_column_letter(c)].width = max(10, len(h) + 4)
    for r, row in enumerate(rows, start=HEADER_ROW + 1):
        for c, v in enumerate(row, start=1):
            ws.cell(r, c, v)
    for i, (label, value) in enumerate(footer or [], start=HEADER_ROW + len(rows) + 2):
        ws.cell(i, 1, label).font = Font(bold=True)
        ws.cell(i, 2, value)
    if footer:
        ws.column_dimensions["A"].width = max(ws.column_dimensions["A"].width or 0, 28)
    ws.freeze_panes = ws.cell(HEADER_ROW + 1, 1)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _table(fmt: Fmt, title: str, note: str, header: list[str], rows: list[list], sheet: str) -> bytes:
    return _csv(header, rows) if fmt == "csv" else _xlsx(title, note, header, rows, sheet)


# ------------------------------------------------------------------- stats


@router.get("/games/{code}/export/stats.{fmt}")
def export_stats(
    code: str,
    fmt: Fmt,
    period: Literal["10", "20", "50", "100", "all"] = "10",
    w10: float = Query(DEFAULT_WEIGHTS[0], ge=0),
    w20: float = Query(DEFAULT_WEIGHTS[1], ge=0),
    w50: float = Query(DEFAULT_WEIGHTS[2], ge=0),
    w100: float = Query(DEFAULT_WEIGHTS[3], ge=0),
    mix: float = Query(50.0, ge=0, le=100),
    session: Session = Depends(get_session),
):
    s = game_stats(code, period, w10, w20, w50, w100, mix, session)
    if s.draw_count == 0:
        raise HTTPException(409, "no draws are stored for this game; import or collect results first")
    header = [
        "number", "appearances", "frequency", "draws_since", "alltime_rank",
        "hot", "cold", "combined", "class",
        "count_10", "count_20", "count_50", "count_100", "count_all",
    ]
    rows = [
        [
            r.number, r.appearances, round(r.frequency, 4), r.gap, r.rank,
            round(r.hot, 4), round(r.cold, 4), round(r.combined, 4), r.cls,
            r.counts["10"], r.counts["20"], r.counts["50"], r.counts["100"], r.counts["all"],
        ]
        for r in s.rows
    ]
    title = f"{code} - number scores ({'all draws' if period == 'all' else 'last ' + period}, {s.draw_count} draws stored)"
    data = _table(fmt, title, NOTE, header, rows, "Scores")
    return _attachment(data, fmt, f"{code}-scores-{period}")


# ------------------------------------------------------------------- draws


@router.get("/games/{code}/export/draws.{fmt}")
def export_draws(code: str, fmt: Fmt, session: Session = Depends(get_session)):
    game = _game(session, code)
    spec = GameSpec.from_model(game)
    draws = session.scalars(
        select(Draw)
        .where(Draw.game_id == game.id)
        .order_by(Draw.draw_date.asc().nulls_first(), Draw.draw_no.asc().nulls_first(), Draw.id.asc())
    ).all()
    if not draws:
        raise HTTPException(409, "no draws are stored for this game; import or collect results first")

    if fmt == "xlsx":
        # The same grid layout the importer reads, so it round-trips (missing numbers stay visible).
        parsed = [
            ParsedDraw(numbers=list(d.numbers), draw_date=d.draw_date, draw_no=d.draw_no, bonus=d.bonus)
            for d in draws
        ]
        data = export_grid(parsed, spec, missing_draw_nos=missing_draw_numbers(session, game))
        return _attachment(data, "xlsx", f"{code}-draws")

    header = ["draw_no", "date"] + [f"n{i}" for i in range(1, spec.k + 1)] + (["bonus"] if spec.has_bonus else [])
    rows = [
        [d.draw_no, d.draw_date.isoformat() if d.draw_date else None, *d.numbers]
        + ([d.bonus] if spec.has_bonus else [])
        for d in draws
    ]
    return _attachment(_csv(header, rows), "csv", f"{code}-draws")


# ------------------------------------------------------------------- combos


class CombosExport(BaseModel):
    method: Literal["hot", "cold", "mixed", "weighted", "random"]
    combinations: list[list[int]] = Field(min_length=1, max_length=MAX_COUNT)


@router.post("/games/{code}/export/combos.{fmt}")
def export_combos(code: str, fmt: Fmt, body: CombosExport, session: Session = Depends(get_session)):
    game = _game(session, code)
    spec = GameSpec.from_model(game)
    combos: list[list[int]] = []
    for i, c in enumerate(body.combinations, start=1):
        try:
            combos.append(validate_numbers(c, spec.n, spec.k))
        except DrawValidationError as exc:
            raise HTTPException(422, f"combination {i}: {exc}") from exc

    header = ["#", "method"] + [f"n{i}" for i in range(1, spec.k + 1)]
    rows = [[i, body.method, *c] for i, c in enumerate(combos, start=1)]
    note = NOTE + " Random picks are the baseline: no method here changes the odds."
    data = _table(fmt, f"{code} - suggested combinations ({body.method})", note, header, rows, "Combinations")
    return _attachment(data, fmt, f"{code}-combinations")


# ------------------------------------------------------------------- bao


@router.get("/games/{code}/export/bao.{fmt}")
def export_bao(
    code: str,
    fmt: Fmt,
    size: int = Query(ge=1, le=18),
    window: int = Query(10, ge=1, le=500),
    end_draw_no: int | None = Query(None, ge=1),
    session: Session = Depends(get_session),
):
    b = bao(code, BaoRequest(size=size, window=window, end_draw_no=end_draw_no), session)
    sure = set(b.sure_in)
    header = ["number", "appearances", "status"]
    rows = [
        [n, b.counts[str(n)], "sure_in" if n in sure else ("filler" if b.unseen_filler else "tie_cut")]
        for n in b.numbers
    ]
    if fmt == "csv":
        return _attachment(_csv(header, rows), "csv", f"{code}-bao-{size}")

    title = f"{code} - Bao {size} (last {window} draws up to draw {b.end_draw_no})"
    summary = [
        ["lines", b.lines],
        [f"ticket price ({b.currency})", b.ticket_price],
        [f"total cost ({b.currency})", b.cost],
        ["jackpot, 1 in", round(b.jackpot_one_in, 1)],
        ["expected hits per draw", round(b.expected_hits, 3)],
    ]
    data = _xlsx(title, b.note, header, rows, "Bao", footer=summary)
    return _attachment(data, "xlsx", f"{code}-bao-{size}")
